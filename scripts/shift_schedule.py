#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ツーマン運行（頭＋助手の2人1組）の月間シフト表を作成する。

制約
  * 1組＝頭1名＋助手1名。平日と土日で出す組数を変えられる。
  * 連続出勤は MAX_CONSECUTIVE 日まで（前月末からの連勤も CARRY_IN_STREAK で引き継ぐ）。
  * 希望休は必ず休みにする。
  * 出勤日数はメンバー間でできるだけ均等にする。
  * 同じ頭と助手の組み合わせが偏らないように毎日ペアを組み替える。
"""

import calendar
import itertools
import random
from collections import defaultdict
from datetime import date

# ---------------------------------------------------------------- 設定 ----
YEAR = 2026
MONTH = 10

# 頭（運転の主担当）と助手。人数が違っても動く。
# 頭は号車が固定（HEADS の並び順に ①②③④…）、助手はその日に組む頭の号車を書く。
# 号車は現行の出勤表と同じ番号（船木①・和田②・永井③・山谷④・横山⑤）
HEADS = ['船木智一', '和田陽向太', '永井海里', '山谷大地', '横山凌']
ASSISTANTS = ['新田', '古川', '山口', '徳留']

# 出勤表に列だけ用意して、ローテーションには入れない人（空欄で出力）
EXTRA_COLUMNS = ['派遣']

CREWS_WEEKDAY = 2       # 平日に出す組数
CREWS_WEEKEND = 4       # 土日に出す組数（土日を多めに）
CREWS_OVERRIDE = {}     # 日にち -> 組数（祝日・繁忙日の個別指定）例: {12: 4}

# 希望休。氏名 -> 日にちのリスト。例: {'頭A': [3, 20], '助手B': [12]}
REQUESTED_OFF = {}

MAX_CONSECUTIVE = 6     # 連続出勤の上限（日）
CARRY_IN_STREAK = {}    # 氏名 -> 前月末時点で何連勤しているか。例: {'頭A': 2}

PREFER_LONG_RUNS = True  # True: 出勤をまとめて連勤気味に / False: 休みを散らす
SEED = 20261001          # 乱数種。変えると別パターンのシフトが出る
TRIES = 400              # 生成の試行回数

OUT = f'{YEAR}年{MONTH}月_出勤表.xlsx'

WEEK_JA = ['月', '火', '水', '木', '金', '土', '日']
CIRCLED = '①②③④⑤⑥⑦⑧⑨⑩'

# ------------------------------------------------------------ 日別の枠 ----


def build_days():
    """その月の各日について、曜日と目標組数を組み立てる。"""
    last = calendar.monthrange(YEAR, MONTH)[1]
    days = []
    for d in range(1, last + 1):
        wd = date(YEAR, MONTH, d).weekday()
        is_weekend = wd >= 5
        target = CREWS_OVERRIDE.get(d, CREWS_WEEKEND if is_weekend else CREWS_WEEKDAY)
        days.append({'day': d, 'wd': wd, 'weekend': is_weekend, 'target': target})
    return days


def off_sets(members):
    return {m: set(REQUESTED_OFF.get(m, [])) for m in members}


def cap_targets(days, heads_off, assist_off):
    """希望休で頭または助手が足りない日は、その日の組数を実際に出せる数まで下げる。"""
    notes = []
    for spec in days:
        d = spec['day']
        avail_h = sum(1 for m in HEADS if d not in heads_off[m])
        avail_a = sum(1 for m in ASSISTANTS if d not in assist_off[m])
        capped = min(spec['target'], avail_h, avail_a)
        if capped < spec['target']:
            notes.append(f'{MONTH}/{d}({WEEK_JA[spec["wd"]]}) 希望休のため {spec["target"]}組 → {capped}組')
        spec['target'] = capped
    return notes


# -------------------------------------------------------------- 割当て ----


def streak_ok(flags, carry):
    """連続出勤が上限以内か。carry は前月末からの連勤日数。"""
    run = carry
    for worked in flags:
        run = run + 1 if worked else 0
        if run > MAX_CONSECUTIVE:
            return False
    return True


def solve_role(members, days, offs, carry_in, rng):
    """1つの職種（頭 or 助手）について、日ごとの出勤者を決める。"""
    n = len(days)
    avail = [[m for m in members if days[i]['day'] not in offs[m]] for i in range(n)]
    # その日に出られる人数と必要人数が同じなら、その人たちは出勤が確定（＝強制出勤）。
    forced = [set(avail[i]) if days[i]['target'] >= len(avail[i]) else set() for i in range(n)]

    def forced_run(start, m):
        """start 日以降、m が連続して強制出勤になる日数。"""
        c = 0
        i = start
        while i < n and m in forced[i]:
            c += 1
            i += 1
        return c

    work = {m: [False] * n for m in members}
    worked = {m: 0 for m in members}
    weekend_worked = {m: 0 for m in members}
    streak = {m: carry_in.get(m, 0) for m in members}
    shortages = []

    for i, spec in enumerate(days):
        need = spec['target']
        cands = []
        for m in avail[i]:
            # 今日出ると、その後の強制出勤日まで含めて上限を超えないか
            if streak[m] + 1 + forced_run(i + 1, m) <= MAX_CONSECUTIVE:
                cands.append(m)
        if len(cands) < need:
            shortages.append((spec['day'], need, len(cands)))
            need = len(cands)

        run = (lambda m: -streak[m]) if PREFER_LONG_RUNS else (lambda m: streak[m])
        if spec['weekend']:
            # 土日は、土日の出勤回数が少ない人から埋める（土日の休みが偏らないように）
            key = lambda m: (weekend_worked[m], worked[m], run(m), rng.random())
        else:
            key = lambda m: (worked[m], run(m), rng.random())
        # 強制出勤の人を先に、残りは出勤日数の少ない人から
        must = [m for m in cands if m in forced[i]]
        rest = sorted((m for m in cands if m not in forced[i]), key=key)
        chosen = (must + rest)[:need]

        for m in members:
            if m in chosen:
                work[m][i] = True
                worked[m] += 1
                if spec['weekend']:
                    weekend_worked[m] += 1
                streak[m] += 1
            else:
                streak[m] = 0

    # 土日と平日を別々に平準化する（土日を先に揃え、そのあと平日で総日数を揃える）
    weekend_idx = [i for i, s in enumerate(days) if s['weekend']]
    weekday_idx = [i for i, s in enumerate(days) if not s['weekend']]
    rebalance(members, days, work, offs, carry_in, rng, weekend_idx)
    rebalance(members, days, work, offs, carry_in, rng, weekday_idx)
    return work, shortages


def rebalance(members, days, work, offs, carry_in, rng, idx_pool):
    """idx_pool の日に限って出勤日を交換し、その範囲の出勤日数を平準化する。"""
    pool = set(idx_pool)
    for _ in range(3000):
        counts = {m: sum(1 for i in pool if work[m][i]) for m in members}
        hi = max(counts, key=lambda m: counts[m])
        lo = min(counts, key=lambda m: counts[m])
        if counts[hi] - counts[lo] <= 1:
            return
        idxs = [i for i in idx_pool
                if work[hi][i] and not work[lo][i] and days[i]['day'] not in offs[lo]]
        rng.shuffle(idxs)
        for i in idxs:
            work[hi][i] = False
            work[lo][i] = True
            if streak_ok(work[hi], carry_in.get(hi, 0)) and streak_ok(work[lo], carry_in.get(lo, 0)):
                break
            work[hi][i] = True
            work[lo][i] = False
        else:
            return


def spread(members, days, offs, carry_in, seed_base):
    """複数回作って、出勤日数のばらつきが最も小さいものを採用する。"""
    best = None
    for t in range(TRIES):
        rng = random.Random(seed_base + t)
        work, shortages = solve_role(members, days, offs, carry_in, rng)
        counts = [sum(work[m]) for m in members]
        we = [sum(1 for i, s in enumerate(days) if s['weekend'] and work[m][i]) for m in members]
        score = (len(shortages), max(we) - min(we), max(counts) - min(counts))
        if best is None or score < best[0]:
            best = (score, work, shortages)
        if score <= (0, 1, 1):
            break
    return best[1], best[2]


# ------------------------------------------------------------ ペア編成 ----


def make_pairs(days, head_work, assist_work, rng):
    """毎日の頭と助手を、過去の組み合わせ回数が少なくなるように組む。"""
    seen = defaultdict(int)
    pairs_by_day = []
    for i, spec in enumerate(days):
        hs = [m for m in HEADS if head_work[m][i]]
        as_ = [m for m in ASSISTANTS if assist_work[m][i]]
        k = min(len(hs), len(as_))
        hs, as_ = hs[:k], as_[:k]
        best, best_cost = None, None
        for perm in itertools.permutations(as_):
            cost = sum(seen[(h, a)] for h, a in zip(hs, perm))
            jitter = rng.random() * 0.01
            if best_cost is None or cost + jitter < best_cost:
                best, best_cost = perm, cost + jitter
        pairs = list(zip(hs, best or ()))
        for h, a in pairs:
            seen[(h, a)] += 1
        pairs_by_day.append(pairs)
    return pairs_by_day, seen


# ---------------------------------------------------------------- 出力 ----


def build_assignment(days, head_work, assist_work, pairs_by_day):
    """各人・各日のセル内容を決める。頭は固定の号車、助手はその日に組む頭の号車。"""
    car_no = {h: i + 1 for i, h in enumerate(HEADS)}
    cells = {m: [''] * len(days) for m in HEADS + ASSISTANTS}
    for i in range(len(days)):
        partner = {a: car_no[h] for h, a in pairs_by_day[i]}
        for m in HEADS:
            cells[m][i] = CIRCLED[car_no[m] - 1] if head_work[m][i] else '休'
        for m in ASSISTANTS:
            if not assist_work[m][i]:
                cells[m][i] = '休'
            else:
                n = partner.get(m)
                cells[m][i] = CIRCLED[n - 1] if n else '○'
    return cells


def max_run(flags, carry):
    run, best = carry, carry
    for w in flags:
        run = run + 1 if w else 0
        best = max(best, run)
    return best


def print_table(days, head_work, assist_work, cells, pairs_by_day, shortages, seen):
    members = HEADS + ASSISTANTS
    width = max(len(m) for m in members) + 2
    hdr = '氏名'.ljust(width) + ' '.join(f'{s["day"]:>2}' for s in days) + '  出勤 土日'
    print(hdr)
    print('-' * len(hdr))
    print('曜日'.ljust(width) + ' '.join(f'{WEEK_JA[s["wd"]]:>2}' for s in days))
    print('車台数'.ljust(width - 1) + ' '.join(f'{len(pairs_by_day[i]):>2}' for i in range(len(days))))
    for label, group, work in (('【頭】', HEADS, head_work), ('【助手】', ASSISTANTS, assist_work)):
        print(label)
        for m in group:
            line = ' '.join(f'{cells[m][i]:>2}' for i in range(len(days)))
            we = sum(1 for i, s in enumerate(days) if s['weekend'] and work[m][i])
            print(m.ljust(width) + line + f'  {sum(work[m]):>3} {we:>3}')
    print()
    print('連勤の最大: ' + ', '.join(
        f'{m}={max_run(work[m], CARRY_IN_STREAK.get(m, 0))}'
        for group, work in ((HEADS, head_work), (ASSISTANTS, assist_work))
        for m in group))
    if shortages:
        print('\n人数が足りず組数を減らした日:')
        for d, want, got in sorted(set(shortages)):
            print(f'  {MONTH}/{d}  {want}組 → {got}組')
    print('\n組み合わせ回数: ' + ', '.join(f'{h}×{a}={c}' for (h, a), c in sorted(seen.items())))


def write_xlsx(days, head_work, assist_work, heads_off, assist_off, cells, pairs_by_day, notes, shortages):
    """手書きの出勤表と同じ様式（縦＝日付、横＝氏名、セル＝号車 or 休）で書き出す。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = '出勤表'

    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center')
    sat_fill = PatternFill('solid', fgColor='DDEBF7')
    sun_fill = PatternFill('solid', fgColor='FCE4E4')
    off_fill = PatternFill('solid', fgColor='F2F2F2')
    req_fill = PatternFill('solid', fgColor='FFF2CC')
    name_fill = PatternFill('solid', fgColor='EFEFEF')
    big = Font(bold=True, size=14)

    members = HEADS + ASSISTANTS
    columns = members + EXTRA_COLUMNS
    ncols = 3 + len(columns)          # 日 / 曜 / 車台数 / 氏名…

    # 見出し行
    ws.cell(1, 3, f'{YEAR}年').font = big
    ws.cell(1, 3).alignment = center
    ws.cell(1, 3 + max(1, len(columns) // 3), '出勤表').font = big
    ws.cell(1, 3 + max(2, len(columns) * 2 // 3), f'{MONTH}月度').font = big

    HEAD_ROW = 2
    for j, title in enumerate(('', '', '車台数')):
        c = ws.cell(HEAD_ROW, 1 + j, title)
        c.font = Font(bold=True)
        c.alignment = center
        c.border = border
    for j, name in enumerate(columns):
        c = ws.cell(HEAD_ROW, 4 + j, name)
        c.font = Font(bold=True)
        c.alignment = center
        c.border = border
        c.fill = name_fill

    for i, spec in enumerate(days):
        r = HEAD_ROW + 1 + i
        c = ws.cell(r, 1, spec['day'])
        c.alignment = center
        c.border = border
        c = ws.cell(r, 2, WEEK_JA[spec['wd']])
        c.alignment = center
        c.border = border
        if spec['wd'] == 5:
            c.font = Font(color='0070C0', bold=True)
            c.fill = sat_fill
        elif spec['wd'] == 6:
            c.font = Font(color='C00000', bold=True)
            c.fill = sun_fill
        c = ws.cell(r, 3, len(pairs_by_day[i]))
        c.alignment = center
        c.border = border
        for j, name in enumerate(columns):
            cell = ws.cell(r, 4 + j)
            cell.alignment = center
            cell.border = border
            if name not in cells:
                continue
            offs = heads_off if name in HEADS else assist_off
            requested = spec['day'] in offs[name]
            cell.value = cells[name][i]
            if cell.value == '休':
                cell.fill = req_fill if requested else off_fill
                if requested:
                    cell.font = Font(bold=True)
            elif spec['wd'] == 5:
                cell.fill = sat_fill
            elif spec['wd'] == 6:
                cell.fill = sun_fill

    # 合計・残業行
    total_row = HEAD_ROW + 1 + len(days)
    ws.cell(total_row, 1, '合計').font = Font(bold=True)
    ws.cell(total_row + 1, 1, '残業').font = Font(bold=True)
    for r in (total_row, total_row + 1):
        for col in range(1, ncols + 1):
            ws.cell(r, col).border = border
            ws.cell(r, col).alignment = center
    ws.cell(total_row, 3, sum(len(p) for p in pairs_by_day))
    for j, name in enumerate(columns):
        if name in cells:
            work = head_work if name in HEADS else assist_work
            ws.cell(total_row, 4 + j, sum(work[name]))

    ws.freeze_panes = 'D3'
    ws.column_dimensions['A'].width = 4.5
    ws.column_dimensions['B'].width = 4.5
    ws.column_dimensions['C'].width = 7
    for j in range(len(columns)):
        ws.column_dimensions[ws.cell(1, 4 + j).column_letter].width = 11

    # --- 組み合わせシート ---
    ws2 = wb.create_sheet('組み合わせ')
    ws2.cell(1, 1, '日').font = Font(bold=True)
    ws2.cell(1, 2, '曜').font = Font(bold=True)
    ws2.cell(1, 3, '車台数').font = Font(bold=True)
    maxc = max((len(p) for p in pairs_by_day), default=0)
    for k in range(maxc):
        ws2.cell(1, 4 + k, f'{CIRCLED[k]}号車').font = Font(bold=True)
    for i, spec in enumerate(days):
        r = 2 + i
        ws2.cell(r, 1, f'{MONTH}/{spec["day"]}').border = border
        c = ws2.cell(r, 2, WEEK_JA[spec['wd']])
        c.border = border
        c.alignment = center
        if spec['wd'] == 5:
            c.fill = sat_fill
        elif spec['wd'] == 6:
            c.fill = sun_fill
        c = ws2.cell(r, 3, len(pairs_by_day[i]))
        c.border = border
        c.alignment = center
        by_car = {HEADS.index(h): (h, a) for h, a in pairs_by_day[i]}
        for k in range(maxc):
            pair = by_car.get(k)
            cell = ws2.cell(r, 4 + k, f'{pair[0]} ／ {pair[1]}' if pair else '')
            cell.border = border
            cell.alignment = center
    ws2.column_dimensions['A'].width = 8
    ws2.column_dimensions['B'].width = 5
    ws2.column_dimensions['C'].width = 7
    for k in range(maxc):
        ws2.column_dimensions[ws2.cell(1, 4 + k).column_letter].width = 24

    # --- 備考シート ---
    ws3 = wb.create_sheet('備考')
    r = 1
    ws3.cell(r, 1, '設定').font = Font(bold=True)
    r += 1
    for k, v in (('対象月', f'{YEAR}年{MONTH}月'),
                 ('平日の組数', CREWS_WEEKDAY),
                 ('土日の組数', CREWS_WEEKEND),
                 ('連続出勤の上限', f'{MAX_CONSECUTIVE}日'),
                 ('頭（号車固定）', '、'.join(f'{CIRCLED[i]}{h}' for i, h in enumerate(HEADS))),
                 ('助手', '、'.join(ASSISTANTS))):
        ws3.cell(r, 1, k)
        ws3.cell(r, 2, v)
        r += 1
    r += 1
    ws3.cell(r, 1, '希望休').font = Font(bold=True)
    r += 1
    for m in members:
        ds = REQUESTED_OFF.get(m)
        if ds:
            ws3.cell(r, 1, m)
            ws3.cell(r, 2, '、'.join(f'{MONTH}/{d}' for d in ds))
            r += 1
    r += 1
    ws3.cell(r, 1, '出勤日数・土日出勤').font = Font(bold=True)
    r += 1
    for group, work in ((HEADS, head_work), (ASSISTANTS, assist_work)):
        for m in group:
            we = sum(1 for i, s in enumerate(days) if s['weekend'] and work[m][i])
            ws3.cell(r, 1, m)
            ws3.cell(r, 2, f'出勤{sum(work[m])}日／土日{we}日／公休{len(days) - sum(work[m])}日／'
                           f'最大{max_run(work[m], CARRY_IN_STREAK.get(m, 0))}連勤')
            r += 1
    if notes or shortages:
        r += 1
        ws3.cell(r, 1, '注意').font = Font(bold=True)
        r += 1
        for note in notes:
            ws3.cell(r, 1, note)
            r += 1
        for d, want, got in sorted(set(shortages)):
            ws3.cell(r, 1, f'{MONTH}/{d} 人数不足のため {want}組 → {got}組')
            r += 1
    ws3.column_dimensions['A'].width = 46
    ws3.column_dimensions['B'].width = 60

    wb.save(OUT)
    return OUT


def verify(days, head_work, assist_work, heads_off, assist_off):
    """希望休・連勤上限・組数が守れているかを確認する。"""
    errors = []
    for members, work, offs in ((HEADS, head_work, heads_off), (ASSISTANTS, assist_work, assist_off)):
        for m in members:
            for i, spec in enumerate(days):
                if work[m][i] and spec['day'] in offs[m]:
                    errors.append(f'{m}: 希望休 {MONTH}/{spec["day"]} に出勤している')
            if max_run(work[m], CARRY_IN_STREAK.get(m, 0)) > MAX_CONSECUTIVE:
                errors.append(f'{m}: 連続出勤が{MAX_CONSECUTIVE}日を超えている')
    for i, spec in enumerate(days):
        h = sum(1 for m in HEADS if head_work[m][i])
        a = sum(1 for m in ASSISTANTS if assist_work[m][i])
        if h != a:
            errors.append(f'{MONTH}/{spec["day"]}: 頭{h}名・助手{a}名で組が作れない')
    return errors


def main():
    rng = random.Random(SEED)
    days = build_days()
    heads_off = off_sets(HEADS)
    assist_off = off_sets(ASSISTANTS)
    notes = cap_targets(days, heads_off, assist_off)

    head_work, sh1 = spread(HEADS, days, heads_off, CARRY_IN_STREAK, SEED)
    assist_work, sh2 = spread(ASSISTANTS, days, assist_off, CARRY_IN_STREAK, SEED + 10000)

    # 頭と助手で出勤人数がずれた日は、少ない方に合わせて組数を確定する
    shortages = sorted(set(sh1) | set(sh2))
    for i, spec in enumerate(days):
        h = sum(1 for m in HEADS if head_work[m][i])
        a = sum(1 for m in ASSISTANTS if assist_work[m][i])
        if h != a:
            more, work, members = (HEADS, head_work, HEADS) if h > a else (ASSISTANTS, assist_work, ASSISTANTS)
            surplus = [m for m in members if work[m][i]]
            surplus.sort(key=lambda m: -sum(work[m]))
            for m in surplus[:abs(h - a)]:
                work[m][i] = False
            shortages.append((spec['day'], spec['target'], min(h, a)))

    pairs_by_day, seen = make_pairs(days, head_work, assist_work, rng)
    cells = build_assignment(days, head_work, assist_work, pairs_by_day)
    print_table(days, head_work, assist_work, cells, pairs_by_day, shortages, seen)
    for note in notes:
        print('注意:', note)

    errors = verify(days, head_work, assist_work, heads_off, assist_off)
    if errors:
        print('\n*** 制約違反 ***')
        for e in errors:
            print(' ', e)
    else:
        print('\n制約チェック: 希望休・連勤上限・組編成すべてOK')

    path = write_xlsx(days, head_work, assist_work, heads_off, assist_off,
                      cells, pairs_by_day, notes, shortages)
    print('出力:', path)


if __name__ == '__main__':
    main()
