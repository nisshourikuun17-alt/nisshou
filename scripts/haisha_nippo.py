# -*- coding: utf-8 -*-
"""配車入力から配車日報（実績帳票）を作る。

入力  配車入力.xlsx     シート『配車入力』『乗務員マスタ』『得意先マスタ』
出力  配車日報_YYYYMMDD.xlsx  シート『配車日報』『乗務員別集計』『得意先別集計』

日報は実際の用紙と同じ左右2ブロック構成にする。
  左ブロック  … 1便目
  右ブロック  … 2便目以降（2便目が1行目、3便目が2行目）
乗務員の並び順は『乗務員マスタ』の表示順に従う。

車番は便ごとに持つ（同じ乗務員でも便で車を替えるため）。ETC高速料金の
車番別集計（etc_allocate_by_vehicle.py）とは車番で突き合わせできる。

    python scripts/haisha_nippo.py --template   入力テンプレートを作る
    python scripts/haisha_nippo.py              入力にある日付ぶんの日報を作る
    python scripts/haisha_nippo.py --date 2026-09-10 --tanto 藁科
"""
import argparse, collections, datetime, os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IN   = os.path.join(ROOT, '配車入力.xlsx')

FONT = '游ゴシック'
WEEK = '月火水木金土日'
BOX  = Border(*[Side(style='thin', color='000000')] * 4)

# 『配車入力』シートの列順。テンプレート作成と読み込みで共用する。
COLS = ['日付', '乗務員', '便', '車番', '積荷', '得意先', '発地', '着地', '備考']

# 日報1ブロックぶんの列と幅。これを左右2つ並べる。
BLOCK = ['乗務員', '車番', '積荷', '得意先', '発地', '着地', '備考']
BWIDTH = (10, 7, 12, 13, 11, 13, 6)

# テンプレートの『記入例』に入れる見本。列順は COLS と同じ。
EXAMPLES = [
    ('2026/9/10', '関', 1, 5022, '玉ねぎ', '鈴与', '日立', '豊洲', ''),
    ('2026/9/10', '関', 2, 5022, '食品', '鈴与', '桶川市', '常陸那珂', ''),
    ('2026/9/10', '芳賀', 1, 1956, 'チーズ', '鈴与', '日立', '阿見', ''),
    ('2026/9/10', '芳賀', 2, 1956, '代用乳', '鈴与', '高崎市', '日立', ''),
    ('2026/9/10', '芳賀', 3, 1956, '紙製品', '鈴与', 'いわき市', '常陸那珂', ''),
    ('2026/9/10', '伊垣', 1, 8816, 'トマト', '豊総合物流', '大洗', '', '着地未定'),
    ('2026/9/10', '', 1, '', '玉葱', '大晴通商', '日立', '豊洲', '乗務員未定'),
]

# テンプレートに入れておくマスタの初期値。実際の日報から起こしたもの。
DRIVERS = [
    ('関', 5022), ('鶴田', 2039), ('伊垣', 8816), ('横須賀', 2185),
    ('仲林', 1464), ('海老澤', 8815), ('芳賀', 1956), ('砂押', 1107),
    ('小林', 8007), ('志賀', 2145), ('宇佐美', 1957), ('萩野間', 1925),
    ('木村', 9480), ('高橋', 5164), ('山崎', 1523), ('根本', 5159),
    ('安野', 3134), ('関根', 1000), ('神長', 3069), ('小堀内', 6891),
    ('吉成', None), ('山崎正二', None), ('石川', None),
]
CUSTOMERS = ['鈴与', 'エアウォーター', '豊総合物流', 'ロードリーム',
             '大晴通商', '光洋運輸', '行方運送', '東亜物産', 'OOCL']


def to_date(v):
    """セルの値を date にする。文字列 'YYYY/M/D' や 'YYYY-MM-DD' も受ける。"""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str) and v.strip():
        y, m, d = (int(x) for x in v.strip().replace('-', '/').split('/'))
        return datetime.date(y, m, d)
    return None


def wareki(d):
    """date -> '令和 8年  9月 10日 （木）'。令和1年 = 2019年。"""
    return '令和 %d年　%d月　%d日　（%s）' % (
        d.year - 2018, d.month, d.day, WEEK[d.weekday()])


def blank(v):
    return v in (None, '')


def read_rows(ws, ncol, start=2):
    """ヘッダ行を飛ばして値だけを返す。全列空欄の行は読み飛ばす。"""
    for row in ws.iter_rows(min_row=start, max_col=ncol, values_only=True):
        if any(not blank(v) for v in row):
            yield list(row) + [None] * (ncol - len(row))


def load(path):
    """入力ブックを読み、明細・乗務員の並び順・得意先マスタを返す。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if '配車入力' not in wb.sheetnames:
        raise SystemExit('『配車入力』シートがありません: %s' % path)

    order, cars = [], {}
    if '乗務員マスタ' in wb.sheetnames:
        for name, car, note in read_rows(wb['乗務員マスタ'], 3):
            if not blank(name):
                order.append(name)
                cars[name] = car

    customers = [c for c, in ((r[0],) for r in
                 read_rows(wb['得意先マスタ'], 1))] if '得意先マスタ' in wb.sheetnames else []

    recs = []
    for i, r in enumerate(read_rows(wb['配車入力'], len(COLS)), 2):
        d = to_date(r[0])
        if d is None:
            raise SystemExit('配車入力 %d行目: 日付が読めません（%r）' % (i, r[0]))
        recs.append(dict(
            date=d, driver=r[1] or '', trip=r[2], car=r[3],
            item=r[4] or '', cust=r[5] or '', from_=r[6] or '',
            to_=r[7] or '', note=r[8] or '', row=i))
    return recs, order, cars, customers


def group(recs, order):
    """乗務員ごとに便をまとめ、マスタの表示順に並べる。

    乗務員が空欄の便は『未配車』としてまとめ、必ず最後に置く。
    マスタに無い乗務員はマスタ掲載分の後ろに、入力に出てきた順で続ける。
    """
    trips = collections.defaultdict(list)
    seen = []
    for r in recs:
        name = r['driver'] if not blank(r['driver']) else None
        if name not in trips:
            seen.append(name)
        trips[name].append(r)

    for v in trips.values():
        # 便No が空欄でも入力順を保てるよう、行番号を第2キーにする
        v.sort(key=lambda r: (r['trip'] if isinstance(r['trip'], int) else 99, r['row']))

    ranked = [n for n in order if n in trips]
    ranked += [n for n in seen if n is not None and n not in order]
    if None in trips:
        ranked.append(None)
    return [(n, trips[n]) for n in ranked]


def style(cell, b=False, sz=10, align=None, fill=None, color=None):
    cell.font = Font(name=FONT, sz=sz, b=b, color=color)
    if align:
        cell.alignment = Alignment(horizontal=align, vertical='center')
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    return cell


def build_nippo(wb, day, recs, order, tanto, spacer):
    """実際の用紙と同じ左右2ブロックの日報シートを作る。"""
    ws = wb.create_sheet('配車日報')
    ncol = len(BLOCK) * 2

    ws['A1'] = '配　車　日　報'
    style(ws['A1'], b=True, sz=18, align='center')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.row_dimensions[1].height = 28

    ws['A2'] = wareki(day)
    style(ws['A2'], b=True, sz=14)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(BLOCK))
    if tanto:
        c = ws.cell(2, len(BLOCK) + 1, '担当　%s　様' % tanto)
        style(c, sz=11, align='right')
        ws.merge_cells(start_row=2, start_column=len(BLOCK) + 1,
                       end_row=2, end_column=ncol)
    ws.row_dimensions[2].height = 22

    for i, h in enumerate(BLOCK * 2, 1):
        c = style(ws.cell(3, i, h), b=True, sz=10, align='center', fill='F2F2F2')
        c.border = BOX
    ws.row_dimensions[3].height = 20

    row = 4
    for name, trips in group(recs, order):
        unassigned = name is None
        head, rest = trips[0], trips[1:]
        # ブロックの高さは「右に並べる2便目以降の数」で決まる。1便だけなら1行。
        height = max(1, len(rest))
        for k in range(height):
            left = head if k == 0 else None
            right = rest[k] if k < len(rest) else None
            for side, rec in ((0, left), (1, right)):
                base = side * len(BLOCK)
                vals = ['' if rec is None else v for v in (
                    name or '未配車', rec and rec['car'], rec and rec['item'],
                    rec and rec['cust'], rec and rec['from_'],
                    rec and rec['to_'], rec and rec['note'])]
                if rec is None:
                    vals = [''] * len(BLOCK)
                for i, v in enumerate(vals, base + 1):
                    c = style(ws.cell(row, i, v if v != '' else None), sz=10)
                    c.border = BOX
                    if i - base in (1, 2, 7):
                        c.alignment = Alignment(horizontal='center', vertical='center')
                    if unassigned and rec is not None:
                        c.fill = PatternFill('solid', fgColor='FCE4E4')
            ws.row_dimensions[row].height = 20
            row += 1
        if spacer:
            for i in range(1, ncol + 1):
                ws.cell(row, i).border = BOX
            ws.row_dimensions[row].height = 14
            row += 1

    n_un = sum(1 for r in recs if blank(r['driver']))
    if n_un:
        style(ws.cell(row, 1, '※未配車が %d 便あります' % n_un),
              b=True, sz=11, color='C00000')

    for i, w in enumerate(BWIDTH * 2, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A4'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1  # 1日分が必ず1ページに収まるようにする
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = '3:3'
    return ws


def build_summary(wb, recs, order, cars):
    """乗務員別・得意先別の便数集計を作る。"""
    ws = wb.create_sheet('乗務員別集計')
    for i, h in enumerate(['乗務員', '車番', '便数', '積荷', '得意先'], 1):
        style(ws.cell(1, i, h), b=True, sz=11, align='center', fill='DDEBF7').border = BOX
    row = 2
    for name, trips in group(recs, order):
        used = sorted({str(t['car']) for t in trips if not blank(t['car'])})
        vals = [name or '未配車', '・'.join(used), len(trips),
                '／'.join(t['item'] for t in trips if t['item']),
                '／'.join(sorted({t['cust'] for t in trips if t['cust']}))]
        for i, v in enumerate(vals, 1):
            style(ws.cell(row, i, v), sz=10)
        row += 1
    style(ws.cell(row, 1, '合計'), b=True, sz=11, align='right')
    style(ws.cell(row, 3, len(recs)), b=True, sz=11)
    for col, w in zip('ABCDE', (12, 14, 7, 26, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'

    ws = wb.create_sheet('得意先別集計')
    for i, h in enumerate(['得意先', '便数', '乗務員'], 1):
        style(ws.cell(1, i, h), b=True, sz=11, align='center', fill='DDEBF7').border = BOX
    cnt = collections.Counter(r['cust'] or '（未記入）' for r in recs)
    who = collections.defaultdict(list)
    for r in recs:
        d = r['driver'] or '未配車'
        if d not in who[r['cust'] or '（未記入）']:
            who[r['cust'] or '（未記入）'].append(d)
    row = 2
    for cust, n in cnt.most_common():
        for i, v in enumerate([cust, n, '・'.join(who[cust])], 1):
            style(ws.cell(row, i, v), sz=10)
        row += 1
    style(ws.cell(row, 1, '合計'), b=True, sz=11, align='right')
    style(ws.cell(row, 2, len(recs)), b=True, sz=11)
    for col, w in zip('ABC', (18, 7, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'


def make_template(path):
    """空の入力ブックを作る。乗務員・得意先マスタは実際の日報から起こした初期値入り。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '配車入力'
    for i, h in enumerate(COLS, 1):
        style(ws.cell(1, i, h), b=True, sz=11, align='center', fill='DDEBF7').border = BOX
    for i, w in enumerate((11, 11, 5, 8, 13, 14, 12, 14, 20), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    ws = wb.create_sheet('乗務員マスタ')
    for i, h in enumerate(['乗務員', '主な車番', '備考'], 1):
        style(ws.cell(1, i, h), b=True, sz=11, align='center', fill='E2EFDA').border = BOX
    for name, car in DRIVERS:
        ws.append([name, car, None])
    for col, w in zip('ABC', (12, 10, 24)):
        ws.column_dimensions[col].width = w
    ws.cell(len(DRIVERS) + 3, 1, '※この並び順が日報の行順になります。')
    ws.cell(len(DRIVERS) + 4, 1, '※車番は目安です。実際の車番は便ごとに配車入力へ。')

    ws = wb.create_sheet('得意先マスタ')
    style(ws.cell(1, 1, '得意先'), b=True, sz=11, align='center', fill='E2EFDA').border = BOX
    for c in CUSTOMERS:
        ws.append([c])
    ws.column_dimensions['A'].width = 18

    # 記入例は『配車入力』に置くと日報に混ざるため、別シートに分ける。
    ws = wb.create_sheet('記入例')
    for i, h in enumerate(COLS, 1):
        style(ws.cell(1, i, h), b=True, sz=11, align='center', fill='FFF2CC').border = BOX
    for r in EXAMPLES:
        ws.append(list(r))
    for i, w in enumerate((11, 11, 5, 8, 13, 14, 12, 14, 20), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for msg in ('', '※このシートは見本です。実際の入力は『配車入力』シートへ。',
                '※便＝1便目・2便目。日報では1便目が左、2便目以降が右に並びます。',
                '※乗務員を空欄にすると日報の最後に「未配車」として赤く出ます。'):
        ws.append([msg])

    wb.save(path)
    print('入力テンプレートを作りました: %s' % path)


def main():
    ap = argparse.ArgumentParser(description='配車入力から配車日報を作る')
    ap.add_argument('--input', default=IN, help='入力ブック（既定: 配車入力.xlsx）')
    ap.add_argument('--date', help='対象日 YYYY-MM-DD（既定: 入力にある全日付）')
    ap.add_argument('--outdir', default=ROOT, help='出力先ディレクトリ')
    ap.add_argument('--tanto', default='', help='担当者名（日報の右上に入る）')
    ap.add_argument('--spacer', action='store_true', help='乗務員ごとに空行を入れる')
    ap.add_argument('--template', action='store_true', help='入力テンプレートを作る')
    args = ap.parse_args()

    if args.template:
        make_template(args.input)
        return

    recs, order, cars, customers = load(args.input)
    if not recs:
        raise SystemExit('『配車入力』シートにデータがありません: %s' % args.input)
    days = sorted({r['date'] for r in recs})
    if args.date:
        want = to_date(args.date)
        if want not in days:
            raise SystemExit('%s のデータが入力にありません' % want)
        days = [want]

    for day in days:
        todays = [r for r in recs if r['date'] == day]
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        build_nippo(wb, day, todays, order, args.tanto, args.spacer)
        build_summary(wb, todays, order, cars)
        out = os.path.join(args.outdir, '配車日報_%s.xlsx' % day.strftime('%Y%m%d'))
        wb.save(out)
        un = sum(1 for r in todays if blank(r['driver']))
        drivers = len({r['driver'] for r in todays if not blank(r['driver'])})
        print('%s  %d便 / 乗務員%d名%s' % (
            day, len(todays), drivers, '  ※未配車 %d便' % un if un else ''))
        print('  -> %s' % out)


if __name__ == '__main__':
    main()
